"""Compare the four supplied workbook exports by type and question order."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1] / "人工智能应用FDE工程师题库"
BOOKS = {
    "原始题目": ROOT / "人工智能应用FDE工程师测试题目 260827.xlsx",
    "学习中心": ROOT / "人工智能应用FDE工程师题库（学习中心）.xlsx",
    "工信人才实名系统": ROOT / "人工智能应用FDE工程师题库（工信人才实名系统）.xlsx",
    "工信人才考试系统": ROOT / "人工智能应用FDE工程师题库（工信人才考试系统）.xlsx",
}
KINDS = {
    "单选题": "single", "单选": "single",
    "多选题": "multiple", "多选": "multiple",
    "判断题": "judgment", "判断": "judgment", "是非题": "judgment",
    "简答题": "short", "问答题": "short",
}


@dataclass(frozen=True)
class Question:
    kind: str
    ordinal: int
    stem: str
    options: tuple[str, ...]
    answer: str
    reference: str
    sheet: str
    row: int


def clean(value: object) -> str:
    return str(value).strip() if value is not None else ""


def normalized(value: str) -> str:
    return re.sub(r"\s+", "", value)


def answer_letters(value: object, kind: str, *, numeric: bool = False) -> str:
    raw = clean(value).upper()
    if kind == "short":
        return ""
    if kind == "judgment" and raw in {"正确", "对", "1"}:
        return "A"
    if kind == "judgment" and raw in {"错误", "错", "2"}:
        return "B"
    if numeric:
        return "".join(sorted(chr(64 + int(char)) for char in raw if char in "12345678"))
    return "".join(sorted(char for char in raw if char in "ABCDEFGH"))


def add(result: dict[str, list[Question]], kind: str, stem: object, options: list[object],
        answer: object, reference: object, sheet: str, row: int, *, numeric: bool = False) -> None:
    if not clean(stem):
        return
    result[kind].append(Question(
        kind, len(result[kind]) + 1, clean(stem),
        tuple(clean(option) for option in options if clean(option)),
        answer_letters(answer, kind, numeric=numeric), clean(reference), sheet, row,
    ))


def read_book(label: str) -> dict[str, list[Question]]:
    workbook = load_workbook(BOOKS[label], read_only=True, data_only=True)
    result: dict[str, list[Question]] = defaultdict(list)
    if label == "原始题目":
        for sheet in ("单选题", "多选题", "判断题", "简答题"):
            kind = KINDS[sheet]
            for row_number, row in enumerate(workbook[sheet].iter_rows(min_row=2, values_only=True), 2):
                if kind == "single":
                    add(result, kind, row[0], list(row[1:5]), row[5], "", sheet, row_number)
                elif kind == "multiple":
                    add(result, kind, row[0], list(row[1:6]), row[7], "", sheet, row_number)
                elif kind == "judgment":
                    add(result, kind, row[0], ["正确", "错误"], row[2], "", sheet, row_number)
                else:
                    add(result, kind, row[0], [], "", row[2], sheet, row_number)
    elif label == "学习中心":
        for row_number, row in enumerate(workbook.active.iter_rows(min_row=2, values_only=True), 2):
            kind = KINDS.get(clean(row[0]))
            if kind:
                options = (list(row[3:10]) if kind in {"single", "multiple"}
                           else ["正确", "错误"] if kind == "judgment" else [])
                add(result, kind, row[2], options, row[10], "", "Sheet1", row_number,
                    numeric=kind != "short")
    elif label == "工信人才实名系统":
        for row_number, row in enumerate(workbook.active.iter_rows(min_row=2, values_only=True), 2):
            kind = KINDS.get(clean(row[1]))
            if kind:
                options = list(row[4:9]) if kind in {"single", "multiple"} else ["正确", "错误"]
                add(result, kind, row[2], options, row[3], "", "Sheet1", row_number)
    else:
        for sheet in ("单选题", "多选题", "判断题", "填空题", "问答题"):
            kind = KINDS.get(sheet)
            if not kind:
                continue
            for row_number, row in enumerate(workbook[sheet].iter_rows(min_row=2, values_only=True), 2):
                if kind in {"single", "multiple"}:
                    add(result, kind, row[0], list(row[1:9]), row[9], "", sheet, row_number)
                elif kind == "judgment":
                    add(result, kind, row[0], ["正确", "错误"], row[1], "", sheet, row_number)
                else:
                    add(result, kind, row[0], [], "", row[1], sheet, row_number)
    workbook.close()
    return result


def audit(*, verbose: bool = True) -> tuple[dict[str, dict[str, list[Question]]], list[tuple[str, Question, dict[str, bool]]]]:
    books = {label: read_book(label) for label in BOOKS}
    original = books["原始题目"]
    differences: list[tuple[str, Question, dict[str, bool]]] = []
    for label, kinds in books.items():
        if verbose:
            counts = {kind: len(rows) for kind, rows in kinds.items()}
            print(f"{label}: {counts}")
        if label == "原始题目":
            continue
        for kind, rows in kinds.items():
            for question in rows:
                if question.ordinal > len(original[kind]):
                    differences.append((label, question, {"extra": True}))
                    continue
                source = original[kind][question.ordinal - 1]
                changed = {
                    "stem": normalized(question.stem) != normalized(source.stem),
                    "options": tuple(map(normalized, question.options)) != tuple(map(normalized, source.options)),
                    "answer": question.answer != source.answer if kind != "short" else False,
                    "reference": normalized(question.reference) != normalized(source.reference)
                    if kind == "short" and question.reference else False,
                }
                if any(changed.values()):
                    differences.append((label, question, changed))
    if verbose:
        for label, question, changed in differences:
            print(f"DIFFERENCE {label} {question.kind} #{question.ordinal} {question.sheet}!{question.row}: "
                  f"{','.join(key for key, value in changed.items() if value)}")
            if changed.get("stem"):
                print(f"  wording: {question.stem}")
    return books, differences


if __name__ == "__main__":
    audit()
