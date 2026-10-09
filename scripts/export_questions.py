"""Export the reviewed FDE source workbook into the static quiz app."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "人工智能应用FDE工程师题库" / "人工智能应用FDE工程师测试题目 260827.xlsx"
DESTINATION = ROOT / "quiz-app" / "questions.json"
SCRIPT_DESTINATION = ROOT / "quiz-app" / "questions.js"


def value(cell: object) -> str:
    return str(cell).strip() if cell is not None else ""


def choices(row: tuple[object, ...], start: int, count: int) -> list[dict[str, str]]:
    return [
        {"key": chr(65 + index), "text": value(row[start + index])}
        for index in range(count)
        if value(row[start + index])
    ]


def export() -> list[dict[str, object]]:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    questions: list[dict[str, object]] = []

    for index, row in enumerate(workbook["单选题"].iter_rows(min_row=2, values_only=True), 1):
        if not value(row[0]):
            continue
        questions.append({
            "id": f"S{index:03d}", "type": "single", "stem": value(row[0]),
            "options": choices(row, 1, 4), "answer": value(row[5]).upper(),
            "explanation": value(row[6]), "chapter": value(row[7]),
        })

    for index, row in enumerate(workbook["多选题"].iter_rows(min_row=2, values_only=True), 1):
        if not value(row[0]):
            continue
        questions.append({
            "id": f"M{index:03d}", "type": "multiple", "stem": value(row[0]),
            "options": choices(row, 1, 5),
            "answer": "".join(char for char in value(row[7]).upper() if char in "ABCDE"),
            "explanation": value(row[9]), "chapter": value(row[10]),
        })

    for index, row in enumerate(workbook["判断题"].iter_rows(min_row=2, values_only=True), 1):
        if not value(row[0]):
            continue
        questions.append({
            "id": f"J{index:03d}", "type": "judgment", "stem": value(row[0]),
            "options": [{"key": "A", "text": "正确"}, {"key": "B", "text": "错误"}],
            "answer": value(row[2]).upper(), "explanation": value(row[1]),
            "chapter": value(row[5]),
        })

    for index, row in enumerate(workbook["简答题"].iter_rows(min_row=2, values_only=True), 1):
        if not value(row[0]):
            continue
        questions.append({
            "id": f"A{index:03d}", "type": "short", "stem": value(row[0]),
            "subquestions": value(row[1]), "reference": value(row[2]),
            "options": [], "answer": "", "explanation": "", "chapter": value(row[3]),
        })

    expected = {"single": 80, "multiple": 40, "judgment": 80, "short": 10}
    counts = {kind: sum(question["type"] == kind for question in questions) for kind in expected}
    if counts != expected or len({question["id"] for question in questions}) != 210:
        raise ValueError(f"Unexpected question counts or duplicate IDs: {counts}")

    for question in questions:
        if not question["stem"] or not question["chapter"]:
            raise ValueError(f"Question or chapter missing: {question['id']}")
        if question["type"] == "short":
            if not question["reference"]:
                raise ValueError(f"Reference answer missing: {question['id']}")
            continue
        option_keys = {option["key"] for option in question["options"]}
        answer = str(question["answer"])
        if not answer or len(set(answer)) != len(answer) or not set(answer) <= option_keys:
            raise ValueError(f"Invalid answer: {question['id']} {answer}")
        if not question["explanation"]:
            raise ValueError(f"Explanation missing: {question['id']}")

    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    DESTINATION.write_text(json.dumps(questions, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    script_data = json.dumps(questions, ensure_ascii=False, separators=(",", ":"))
    SCRIPT_DESTINATION.write_text(f"window.FDE_QUESTIONS = {script_data};\n", encoding="utf-8")
    print(f"Exported {len(questions)} questions: {counts}")
    return questions


if __name__ == "__main__":
    export()
